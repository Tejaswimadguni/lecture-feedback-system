import os
import json
import argparse
from typing import Dict, Any, Optional
from pptx import Presentation


def derive_lecture_id(file_path: str) -> str:
    """Derives a lecture ID from a filename by stripping path and extension."""
    basename = os.path.basename(file_path)
    lecture_id, _ = os.path.splitext(basename)
    return lecture_id


def extract_ppt_text(ppt_path: str, lecture_id: Optional[str] = None) -> Dict[str, Any]:
    """
    Extracts text from a PowerPoint (.pptx) file slide by slide.

    Args:
        ppt_path: Path to the input PPTX file.
        lecture_id: Identifier for the lecture. If None, derived from filename.

    Returns:
        Dict containing lecture_id and a list of page objects with 1-indexed slide numbers and text.
    """
    if not os.path.exists(ppt_path):
        raise FileNotFoundError(f"PPT file not found at: {ppt_path}")

    if not lecture_id:
        lecture_id = derive_lecture_id(ppt_path)

    prs = Presentation(ppt_path)
    pages = []

    for idx, slide in enumerate(prs.slides, start=1):
        slide_text_pieces = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        slide_text_pieces.append(text)
        
        full_slide_text = "\n".join(slide_text_pieces).strip()
        pages.append({
            "page": idx,
            "text": full_slide_text
        })

    return {
        "lecture_id": lecture_id,
        "pages": pages
    }


def save_ppt_text_json(ppt_path: str, output_path: Optional[str] = None, lecture_id: Optional[str] = None) -> str:
    """
    Extracts text from PPTX and saves it to a JSON file.

    Args:
        ppt_path: Path to the input PPTX file.
        output_path: Path where JSON should be saved. If None, defaults to data/processed/pdf_text/<lecture_id>_slides.json.
        lecture_id: Identifier for the lecture. If None, derived from filename.

    Returns:
        Path to the saved JSON file.
    """
    data = extract_ppt_text(ppt_path, lecture_id=lecture_id)

    if not output_path:
        base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
        output_dir = os.path.join(base_dir, "data", "processed", "pdf_text")
        os.makedirs(output_dir, exist_ok=True)
        output_path = os.path.join(output_dir, f"{data['lecture_id']}_slides.json")
    else:
        parent_dir = os.path.dirname(os.path.abspath(output_path))
        if parent_dir:
            os.makedirs(parent_dir, exist_ok=True)

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

    return output_path


def main():
    parser = argparse.ArgumentParser(description="Ingest PowerPoint presentation slides into structured JSON.")
    parser.add_argument("ppt_path", help="Path to input PPTX file")
    parser.add_argument("--output", "-o", help="Optional path to output JSON file")
    parser.add_argument("--lecture-id", "-l", help="Optional lecture ID")

    args = parser.parse_args()

    saved_path = save_ppt_text_json(args.ppt_path, output_path=args.output, lecture_id=args.lecture_id)
    print(f"Successfully processed PPT into JSON: {saved_path}")


if __name__ == "__main__":
    main()